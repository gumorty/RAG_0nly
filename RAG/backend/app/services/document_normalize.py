import csv
import io
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import chardet
from bs4 import BeautifulSoup
from markdownify import markdownify as html_to_markdown
from openpyxl import load_workbook
from pptx import Presentation


CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".cpp", ".c", ".h", ".hpp",
    ".cs", ".php", ".rb", ".swift", ".kt", ".scala", ".sql", ".sh", ".ps1", ".bat", ".yaml",
    ".yml", ".toml", ".ini", ".conf", ".json", ".xml", ".css", ".scss", ".html", ".vue",
}


def normalize_for_ragflow(filename: str, data: bytes, content_type: str | None) -> tuple[str, bytes, str | None, dict]:
    """Create a text-friendly shadow file for formats that often fail native parsing.

    The original file is still stored by the management platform. RAGFlow receives
    a Markdown/CSV-like representation so parsing, chunking, embedding and retrieval
    remain stable across legacy spreadsheets, HTML exports and slide decks.
    """
    ext = Path(filename).suffix.lower()
    try:
        if ext in {".xlsx", ".xlsm"}:
            text = _xlsx_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "openpyxl"}
        if ext == ".csv":
            text = _csv_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "csv"}
        if ext == ".xls":
            text = _xls_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "xlrd"}
        if ext == ".docx":
            text = _docx_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "python-docx"}
        if ext == ".doc":
            text = _doc_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "antiword"}
        if ext in {".html", ".htm"}:
            text = _html_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "beautifulsoup"}
        if ext == ".pptx":
            text = _pptx_to_markdown(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "python-pptx"}
        if ext in {".md", ".markdown"}:
            text = _decode_text(data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "markdown"}
        if ext in {".txt", ".log", ".rst"}:
            text = _plain_text_to_markdown(filename, data)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "plain-text"}
        if ext in CODE_EXTENSIONS:
            text = _code_to_markdown(filename, data, ext)
            return _shadow(filename), text.encode("utf-8"), "text/markdown", {"normalized_for_ragflow": True, "normalizer": "code"}
    except Exception as exc:
        return filename, data, content_type, {"normalization_warning": str(exc)}
    return filename, data, content_type, {}


def _shadow(filename: str) -> str:
    return f"{filename}.ragflow.md"


def _xlsx_to_markdown(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}\n")
        rows = [[_cell(cell) for cell in row] for row in ws.iter_rows(values_only=True)]
        parts.append(_rows_to_markdown(rows))
    return "\n\n".join(part for part in parts if part.strip())


def _docx_to_markdown(data: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("python-docx is required to normalize .docx files") from exc

    doc = DocxDocument(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        if "heading 1" in style or "标题 1" in style:
            parts.append(f"# {text}")
        elif "heading 2" in style or "标题 2" in style:
            parts.append(f"## {text}")
        elif "heading" in style or "标题" in style:
            parts.append(f"### {text}")
        else:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        rows = [[_cell(cell.text) for cell in row.cells] for row in table.rows]
        table_text = _rows_to_markdown(rows)
        if table_text:
            parts.append(f"## Word Table {table_index}\n\n{table_text}")
    return "\n\n".join(part for part in parts if part.strip())


def _doc_to_markdown(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        result = subprocess.run(
            ["antiword", tmp_path],
            check=True,
            capture_output=True,
            timeout=60,
        )
        text = result.stdout.decode("utf-8", errors="replace").strip()
        return f"# Legacy Word document\n\n{text}"
    finally:
        try:
            Path(tmp_path).unlink(missing_ok=True)
        except OSError:
            pass


def _xls_to_markdown(data: bytes) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("xlrd is required to normalize legacy .xls files") from exc
    book = xlrd.open_workbook(file_contents=data)
    parts = []
    for sheet in book.sheets():
        parts.append(f"# Sheet: {sheet.name}\n")
        rows = [[_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        parts.append(_rows_to_markdown(rows))
    return "\n\n".join(part for part in parts if part.strip())


def _csv_to_markdown(data: bytes) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_markdown(rows)


def _html_to_markdown(data: bytes) -> str:
    html = _decode_text(data)
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.string.strip() if soup.title and soup.title.string else "HTML document"
    markdown = html_to_markdown(str(soup), heading_style="ATX")
    return f"# {title}\n\n{markdown}"


def _pptx_to_markdown(data: bytes) -> str:
    deck = Presentation(io.BytesIO(data))
    parts = []
    for index, slide in enumerate(deck.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append([_cell(cell.text) for cell in row.cells])
                table_text = _rows_to_markdown(rows)
                if table_text:
                    lines.append(table_text)
                continue
            text = getattr(shape, "text", "")
            if text and text.strip():
                lines.append(text.strip())
        if lines:
            parts.append(f"# Slide {index}\n\n" + "\n\n".join(lines))
    return "\n\n".join(parts)


def _plain_text_to_markdown(filename: str, data: bytes) -> str:
    return f"# {filename}\n\n{_decode_text(data)}"


def _code_to_markdown(filename: str, data: bytes, ext: str) -> str:
    language = ext.lstrip(".") or "text"
    text = _decode_text(data)
    return f"# Code: {filename}\n\n```{language}\n{text}\n```"


def _decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig", errors="replace")
    detected = chardet.detect(data[:200_000])
    encoding = detected.get("encoding") or "utf-8"
    return data.decode(encoding, errors="replace")


def _rows_to_markdown(rows: list[list[str]], max_rows: int = 2000) -> str:
    clean_rows = [_trim_empty(row) for row in rows[:max_rows]]
    clean_rows = [row for row in clean_rows if any(cell.strip() for cell in row)]
    if not clean_rows:
        return ""
    width = max(len(row) for row in clean_rows)
    normalized = [row + [""] * (width - len(row)) for row in clean_rows]
    header = normalized[0]
    body = normalized[1:]
    lines = [
        "| " + " | ".join(_escape(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(_escape(cell) for cell in row) + " |")
    if len(rows) > max_rows:
        lines.append(f"\n\n> 表格行数超过 {max_rows}，已截取前 {max_rows} 行进入索引。")
    return "\n".join(lines)


def _trim_empty(row: list[str]) -> list[str]:
    items = list(row)
    while items and not items[-1].strip():
        items.pop()
    return items


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _escape(value: str) -> str:
    return value.replace("|", "\\|")
