from docx import Document as DocxDocument

from app.services.document_normalize import normalize_for_ragflow


def test_normalize_docx_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "report.docx"
    doc = DocxDocument()
    doc.add_heading("Enterprise RAG Report", level=1)
    doc.add_paragraph("This document describes parsing quality.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Recall"
    table.cell(1, 1).text = "0.92"
    doc.save(path)

    name, data, content_type, meta = normalize_for_ragflow(path.name, path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    text = data.decode("utf-8")

    assert name == "report.docx.ragflow.md"
    assert content_type == "text/markdown"
    assert "# Enterprise RAG Report" in text
    assert "Metric" in text and "Recall" in text
    assert meta["normalizer"] == "python-docx"


def test_normalize_code_wraps_fenced_block():
    name, data, content_type, meta = normalize_for_ragflow("main.py", b"print('hello')\n", "text/x-python")
    text = data.decode("utf-8")

    assert name == "main.py.ragflow.md"
    assert content_type == "text/markdown"
    assert "```py" in text
    assert "print('hello')" in text
    assert meta["normalizer"] == "code"
