from pathlib import Path
import subprocess

import chardet
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from markdownify import markdownify
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.rag.normalize import normalize_text
from app.rag.schemas import ParsedDocument


class DocumentParser:
    def parse(self, path: Path, filename: str, content_type: str | None = None) -> ParsedDocument:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = self._parse_pdf(path)
        elif suffix == ".docx":
            text = self._parse_docx(path)
        elif suffix == ".doc":
            text = self._parse_doc(path)
        elif suffix == ".pptx":
            text = self._parse_pptx(path)
        elif suffix in {".xlsx", ".xlsm"}:
            text = self._parse_xlsx(path)
        elif suffix in {".html", ".htm"}:
            text = self._parse_html(path)
        elif suffix in {".md", ".markdown", ".txt", ".csv", ".json"}:
            text = self._parse_text(path)
        else:
            raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

        text = normalize_text(text)
        if not text:
            raise ValueError("No readable text was extracted from this document.")
        return ParsedDocument(
            title=Path(filename).stem,
            text=text,
            sections=self._extract_sections(text),
            metadata={"filename": filename, "content_type": content_type, "suffix": suffix},
        )

    def _parse_pdf(self, path: Path) -> str:
        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages.append(f"\n\n[Page {index}]\n{page_text}")
        return "\n".join(pages)

    def _parse_docx(self, path: Path) -> str:
        doc = DocxDocument(str(path))
        lines = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    def _parse_doc(self, path: Path) -> str:
        result = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(path)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            timeout=60,
        )
        text = result.stdout.strip()
        if result.returncode != 0 or not text:
            error = (result.stderr or "antiword did not extract text").strip()
            raise ValueError(f"Failed to parse legacy .doc file: {error}")
        return text

    def _parse_pptx(self, path: Path) -> str:
        presentation = Presentation(str(path))
        lines = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines.append(f"\n\n[Slide {index}]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        return "\n".join(lines)

    def _parse_xlsx(self, path: Path) -> str:
        workbook = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"\n\n[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    def _parse_html(self, path: Path) -> str:
        raw = path.read_bytes()
        html = _decode_text_bytes(raw)
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.extract()
        return markdownify(str(soup))

    def _parse_text(self, path: Path) -> str:
        return _decode_text_bytes(path.read_bytes())

    def _extract_sections(self, text: str) -> list[dict]:
        sections: list[dict] = []
        current_title = "正文"
        current_lines: list[str] = []
        for line in text.splitlines():
            stripped = line.strip()
            is_heading = stripped.startswith("#") or (
                len(stripped) <= 80 and stripped.endswith(("：", ":")) and len(stripped) > 2
            )
            if is_heading and current_lines:
                sections.append({"title_path": [current_title], "text": "\n".join(current_lines).strip()})
                current_lines = []
                current_title = stripped.lstrip("#").strip(" :：") or current_title
            else:
                current_lines.append(line)
        if current_lines:
            sections.append({"title_path": [current_title], "text": "\n".join(current_lines).strip()})
        return [section for section in sections if section["text"]]


def _decode_text_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    detected = chardet.detect(raw).get("encoding") or "utf-8"
    return raw.decode(detected, errors="ignore")
